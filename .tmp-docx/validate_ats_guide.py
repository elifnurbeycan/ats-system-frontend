import re
import zipfile

from docx import Document


path = r"C:\Users\elifb\Desktop\manus-ats-frontend\docs\ATS-Yerel-Kurulum-ve-Veritabani-Yedekleme-Rehberi.docx"
document = Document(path)
headings = [p.text for p in document.paragraphs if p.style.name.startswith("Heading")]

with zipfile.ZipFile(path) as archive:
    xml = archive.read("word/document.xml").decode("utf-8")
    relationships = archive.read("word/_rels/document.xml.rels").decode("utf-8")

print(f"headings={len(headings)}")
print("\n".join(headings))
print(f"bookmarks={len(re.findall(r'<w:bookmarkStart', xml))}")
print(f"internal_links={len(re.findall(r'w:anchor=', xml))}")
print(f"external_links={len(re.findall(r'TargetMode=\"External\"', relationships))}")
print(f"tables={len(document.tables)}")
print(f"paragraphs={len(document.paragraphs)}")
anchors = set(re.findall(r'w:anchor="([^"]+)"', xml))
bookmark_names = set(re.findall(r'w:bookmarkStart[^>]+w:name="([^"]+)"', xml))
print(f"unmatched_internal_links={sorted(anchors - bookmark_names)}")
print(f"native_toc_field={'TOC ' in xml and 'fldCharType=\"begin\"' in xml}")
settings_xml = zipfile.ZipFile(path).read("word/settings.xml").decode("utf-8")
print(f"update_fields_on_open={'updateFields' in settings_xml}")
text = "\n".join(p.text for p in document.paragraphs)
print(f"frontend_3000={'http://localhost:3000' in text}")
print(f"backend_8080={'http://localhost:8080' in text}")
print(f"postgres_5432={'5432' in text}")
print(f"repo_seed_dump={'database\\ats_system_seed.dump' in text}")
